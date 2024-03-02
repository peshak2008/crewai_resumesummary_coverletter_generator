from docx import Document
import os
from utils.config import (
    resume_template, resume_sections_to_update, scrape_config, resume_name,
    notion_config, candidate_name, default_job_title, default_skills_summary,
    default_location, llm_model, create_pdf
)
import subprocess

class ErrorChecker:
    def __init__(self):
        self.default_resume_folder = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'default_resume')

    def check_config_values(self):
        # Check for essential configuration values
        config_values = {
            "candidate_name": candidate_name,
            "default_job_title": default_job_title,
            "default_skills_summary": default_skills_summary,
            "default_location": default_location,
            "resume_name": resume_name,
            "resume_template": resume_template,
            "llm_model": llm_model,
        }
        missing_configs = [key for key, value in config_values.items() if not value]
        if missing_configs:
            raise ValueError(f"Missing configuration values in config.py: {', '.join(missing_configs)}")

    def check_libreoffice_installation(self):
        if create_pdf:  # Only proceed if create_pdf is True
            try:
                # Attempt to run the soffice command with --version to check if LibreOffice is installed
                subprocess.run(['soffice', '--version'], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                print("LibreOffice is installed and accessible.")
            except (subprocess.CalledProcessError, FileNotFoundError):
                raise EnvironmentError("LibreOffice is not installed or not accessible from the command line. Please ensure LibreOffice is installed and 'soffice' command can be run.")
        else:
            print("PDF creation is disabled. Skipping LibreOffice installation check.")


    def check_template_resume(self):
        template_path = os.path.join(self.default_resume_folder, resume_template)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template resume file does not exist at path: {template_path}")
        
        doc = Document(template_path)
        template_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        
        required_placeholders = {
            'title': '{TITLE}',
            'summary': '{SUMMARY}',
            'location': '{LOCATION}',
        }
        
        missing_placeholders = []
        for section, placeholder in required_placeholders.items():
            if resume_sections_to_update.get(section, False) and placeholder not in template_content:
                missing_placeholders.append(placeholder)
        
        if missing_placeholders:
            raise ValueError(f"Template resume is missing placeholders for: {', '.join(missing_placeholders)}")

    def check_resume_file(self):
        resume_path = os.path.join(self.default_resume_folder, resume_name)
        if not os.path.exists(resume_path):
            raise FileNotFoundError(f"Resume file does not exist at path: {resume_path}")

    def check_output_json(self):
        if not scrape_config.get("perform_scrape", False):
            output_file_path = os.path.join(os.path.dirname(self.default_resume_folder), 'output.json')
            if not os.path.exists(output_file_path):
                raise FileNotFoundError("output.json does not exist. Please enable scraping or manually provide the file in the root directory.")

    def check_notion_config(self):
        if notion_config.get("save_to_notion", False):
            if "NOTION_INTEGRATION_TOKEN" not in os.environ:
                raise EnvironmentError("NOTION_INTEGRATION_TOKEN is not set in the environment variables.")
            if "database_id" not in notion_config or not notion_config["database_id"]:
                raise ValueError("database_id is not provided in the notion_config.")
        else:
            print("Notion integration is disabled.")

    def run_checks(self):
        self.check_config_values()
        self.check_notion_config()
        self.check_libreoffice_installation()
        self.check_template_resume()
        self.check_resume_file()
        self.check_output_json()
        print("All checks passed. Tool will now run.")

def check_errors():
    checker = ErrorChecker()
    checker.run_checks()

if __name__ == "__main__":
    check_errors()